import openpyxl

wb = openpyxl.load_workbook("sts thyroid data.xlsx")

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    comments = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.comment:
                comments.append((cell.coordinate, cell.comment.text))
    print(f"Sheet {sheetname} comments ({len(comments)}): {comments}")
