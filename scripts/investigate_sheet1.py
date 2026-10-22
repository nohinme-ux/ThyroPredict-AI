import pandas as pd
import openpyxl

wb = openpyxl.load_workbook("sts thyroid data.xlsx", data_only=True)
ws1 = wb["Sheet1"]
headers1 = [ws1.cell(1, c).value for c in range(1, ws1.max_column + 1)]
print("Sheet 1 headers:")
for idx, h in enumerate(headers1):
    if h is not None:
        print(f"  col {idx+1}: {h}")

# Check what is in column 21 (which was Unnamed: 20 in pandas)
print("\nValues in Sheet 1, column 21:")
col21_vals = [ws1.cell(r, 21).value for r in range(1, ws1.max_row + 1)]
print("Header of col 21:", ws1.cell(1, 21).value)
non_empty_c21 = [(r, ws1.cell(r, 21).value) for r in range(1, ws1.max_row + 1) if ws1.cell(r, 21).value is not None]
print(f"Non-empty in col 21 ({len(non_empty_c21)}):", non_empty_c21[:20])

# Check all cells in sheet1 to see if there is any other column with labels
print("\nChecking any text values in Sheet 1 across all columns:")
for c in range(1, ws1.max_column + 1):
    col_name = ws1.cell(1, c).value
    col_vals = [ws1.cell(r, c).value for r in range(2, ws1.max_row + 1)]
    str_vals = [str(v).strip() for v in col_vals if v is not None and not isinstance(v, (int, float))]
    unique_strs = set(str_vals)
    if unique_strs and len(unique_strs) < 20:
        print(f"Col {c} ({col_name}): {unique_strs}")
