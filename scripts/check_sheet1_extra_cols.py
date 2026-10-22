import openpyxl

wb = openpyxl.load_workbook("sts thyroid data.xlsx", data_only=True)
ws1 = wb["Sheet1"]

print("Sheet 1 max_col:", ws1.max_column)
for c in range(48, ws1.max_column + 1):
    header = ws1.cell(1, c).value
    col_vals = [ws1.cell(r, c).value for r in range(2, 25)]
    non_none = [v for v in col_vals if v is not None]
    print(f"Col {c}: Header='{header}', sample non-none={non_none[:5]}")
